"""Which office carries a trade, and what follows from it.

A trade belongs to Beirut or to Cairo; a deliverable is worked by whichever
offices its trades belong to, in the proportions those trades carry. Everything
here is that one sentence, checked: that the split follows the allocations
rather than being typed anywhere, that the figures add up per office, and that
a trade nobody has placed reads as unplaced rather than as Beirut.
"""

from __future__ import annotations

from app import offices


def text(response) -> str:
    return response.get_data(as_text=True)


# --- the office itself ------------------------------------------------------

def test_an_office_is_recognised_however_it_is_typed():
    assert offices.normalise("Cairo") == "cairo"
    assert offices.normalise("CAIRO") == "cairo"
    assert offices.normalise("beirut") == "beirut"
    assert offices.normalise(" Beirut ") == "beirut"


def test_anything_that_is_not_an_office_is_blank_rather_than_a_guess():
    for value in ("Dubai", "", None, 7, "bei"):
        assert offices.normalise(value) == ""


def test_a_trade_nobody_has_placed_reads_as_unplaced():
    assert offices.name_of("") == "Not set"
    assert offices.name_of("cairo") == "Cairo"


# --- what a deliverable inherits -------------------------------------------

def test_a_deliverable_is_worked_by_the_offices_its_trades_belong_to():
    where = offices.by_id([{"id": 1, "office": "beirut"}, {"id": 2, "office": "cairo"}])
    task = {"allocations": {1: 0.6, 2: 0.4}}

    assert offices.shares(task, where) == {"beirut": 0.6, "cairo": 0.4}
    # The larger share first, so a row reads as the office mostly doing it.
    assert offices.of_task(task, where) == ["beirut", "cairo"]
    assert offices.label(offices.of_task(task, where)) == "Beirut + Cairo"


def test_two_trades_in_one_office_are_one_share_not_two():
    where = offices.by_id([{"id": 1, "office": "cairo"}, {"id": 2, "office": "cairo"}])
    assert offices.shares({"allocations": {1: 0.3, 2: 0.7}}, where) == {"cairo": 1.0}


def test_a_share_carried_by_an_unplaced_trade_is_not_counted_as_an_office():
    where = offices.by_id([{"id": 1, "office": "beirut"}, {"id": 2, "office": ""}])
    task = {"allocations": {1: 0.5, 2: 0.5}}
    assert offices.of_task(task, where) == ["beirut"]
    assert offices.shares(task, where)[""] == 0.5


# --- the rollup -------------------------------------------------------------

def test_the_figures_add_up_per_office():
    rows = [
        {"name": "Marine", "office": "beirut", "scope_weight_pct": 0.5,
         "earned_contribution": 0.25, "planned_contribution": 0.3,
         "budget_hours": 1000, "spent_hours": 600, "earned_hours": 500},
        {"name": "Utilities", "office": "cairo", "scope_weight_pct": 0.5,
         "earned_contribution": 0.1, "planned_contribution": 0.1,
         "budget_hours": 200, "spent_hours": 50, "earned_hours": 40},
    ]
    beirut, cairo = offices.rollup(rows, hours_per_month=176)

    assert beirut["name"] == "Beirut" and cairo["name"] == "Cairo"
    assert beirut["budget_hours"] == 1000
    # Half the scope, half of it earned: 50% through what Beirut is carrying.
    assert abs(beirut["earned_pct_of_office"] - 0.5) < 1e-9
    assert abs(beirut["planned_pct_of_office"] - 0.6) < 1e-9
    assert abs(beirut["cpi"] - 500 / 600) < 1e-9
    assert cairo["trades"] == ["Utilities"]


def test_unplaced_trades_are_rolled_up_last_rather_than_dropped():
    rows = [
        {"name": "Odd job", "office": "", "scope_weight_pct": 0.1, "earned_contribution": 0,
         "planned_contribution": 0, "budget_hours": 10, "spent_hours": 0, "earned_hours": 0},
        {"name": "Marine", "office": "beirut", "scope_weight_pct": 0.9, "earned_contribution": 0,
         "planned_contribution": 0, "budget_hours": 90, "spent_hours": 0, "earned_hours": 0},
    ]
    rolled = offices.rollup(rows)
    assert [r["name"] for r in rolled] == ["Beirut", "Not set"]


# --- on the project ---------------------------------------------------------

def test_the_snapshot_carries_the_office_on_every_trade_and_deliverable(app):
    from app.db import query_one
    from app.service import as_dict, project_snapshot

    with app.app_context():
        snapshot = project_snapshot(as_dict(query_one("SELECT * FROM projects WHERE id = 1")))

    assert {t["office"] for t in snapshot["trades"]} == {"beirut", "cairo"}
    assert [o["name"] for o in snapshot["offices"]] == ["Beirut", "Cairo"]
    assert all(task["offices"] for task in snapshot["tasks"]), \
        "every seeded deliverable has a trade, so every one has an office"

    # The offices carry the whole scope between them.
    total = sum(o["scope_weight_pct"] for o in snapshot["offices"])
    assert abs(total - 1) < 0.005


# --- on the screens ---------------------------------------------------------

def test_the_setup_sheet_offers_the_office_against_each_trade(signed_in):
    page = text(signed_in.get("/projects/1/setup?password=2026"))
    assert 'name="trade_' in page and "_office" in page
    assert ">Beirut<" in page and ">Cairo<" in page


def test_an_office_is_saved_from_the_setup_sheet(signed_in, app):
    from app.db import query_one

    signed_in.post("/projects/1/setup/unlock", data={"password": "2026"})
    with app.app_context():
        trade = query_one("SELECT id FROM trades WHERE project_id = 1 AND name = 'Utilities'")
    signed_in.post(f"/projects/1/trades/{trade['id']}",
                   data={"name": "Utilities", "budget_hours": "88", "office": "Beirut",
                         "color": "#eda100"}, follow_redirects=True)
    with app.app_context():
        after = query_one("SELECT office FROM trades WHERE id = ?", (trade["id"],))
    assert after["office"] == "beirut"


def test_progress_can_be_filtered_to_one_office(signed_in):
    everything = text(signed_in.get("/projects/1/tasks"))
    cairo = text(signed_in.get("/projects/1/tasks?office=cairo"))

    assert "55 shown" in everything
    assert "55 shown" not in cairo
    # The chips are there to change it back.
    assert "office=cairo" in everything


def test_an_office_nobody_has_is_not_a_filter_that_hides_everything(signed_in):
    """A stray ?office=dubai reads as no filter rather than as an empty page."""
    page = text(signed_in.get("/projects/1/tasks?office=dubai"))
    assert "55 shown" in page


def test_the_budget_adds_the_hours_up_per_office(signed_in):
    page = text(signed_in.get("/projects/1/budget"))
    assert "By office" in page
    assert "Beirut" in page and "Cairo" in page


def test_the_office_goes_out_to_excel_and_comes_back(app, signed_in):
    import io

    from openpyxl import load_workbook

    answer = signed_in.get("/projects/1/setup/export")
    book = load_workbook(io.BytesIO(answer.data))
    rows = {r[0]: r for r in book["Trades"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert rows["Marine"][3] == "Beirut"
    assert rows["Geotechnical"][3] == "Cairo"

    from app.excel import read_workbook

    parsed = read_workbook(answer.data)
    assert {t["name"]: t["office"] for t in parsed["trades"]}["Marine"] == "beirut"
